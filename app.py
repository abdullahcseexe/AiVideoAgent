import os
import io
import html
import re
import tempfile
from datetime import datetime
from urllib.parse import urlparse

import streamlit as st
from docx import Document
from dotenv import load_dotenv

from utils.audio_processor import process_input, cleanup
from core.transcriber import transcribe_all
from core.summarizer import summarize, generate_title
from core.extractor import extract_action_items, extract_key_decisions, extract_questions
from core.rag_engine import build_rag_chain, ask_question

load_dotenv()

st.set_page_config(
    page_title="Video Assistant",
    page_icon="◼",
    layout="centered",
    initial_sidebar_state="collapsed",
)

# ---------------------------------------------------------------- styling
st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600&family=Source+Serif+4:wght@500;600&display=swap');

    :root {
        --ink: #1F2933;
        --muted: #667380;
        --line: #DADFE3;
        --surface: #FFFFFF;
        --accent: #3B4A5A;
    }

    html, body, [class*="css"], .stMarkdown, .stTextInput, .stButton, .stTabs, .stRadio {
        font-family: 'Inter', -apple-system, 'Segoe UI', sans-serif;
        color: var(--ink);
    }

    /* this app has no sidebar */
    section[data-testid="stSidebar"],
    [data-testid="stSidebarCollapsedControl"],
    [data-testid="collapsedControl"],
    [data-testid="stExpandSidebarButton"] { display: none !important; }

    #MainMenu, footer, header[data-testid="stHeader"] { visibility: hidden; height: 0; }
    .block-container { padding-top: 2.5rem; padding-bottom: 4rem; max-width: 900px; }

    h1, h2, h3 { font-family: 'Source Serif 4', Georgia, serif; font-weight: 600; letter-spacing: -0.01em; }
    h1 { font-size: 2.1rem; line-height: 1.25; }

    .brand { font-family: 'Source Serif 4', Georgia, serif; font-size: 1.3rem; font-weight: 600; margin: 0; }
    .brand-sub { color: var(--muted); font-size: 0.9rem; margin: 0.15rem 0 0 0; }

    /* buttons */
    .stButton > button, .stDownloadButton > button {
        border-radius: 4px; border: 1px solid var(--line); background: var(--surface);
        color: var(--ink); font-weight: 500; padding: 0.55rem 1rem; transition: border-color .15s, background .15s;
    }
    .stButton > button:hover, .stDownloadButton > button:hover { border-color: var(--accent); color: var(--accent); background: var(--surface); }
    .stButton > button[kind="primary"] { background: var(--accent); border-color: var(--accent); color: #fff; }
    .stButton > button[kind="primary"]:hover { background: #2C3947; border-color: #2C3947; color: #fff; }
    .stButton > button:disabled { opacity: 0.5; }

    /* inputs */
    .stTextInput input { border-radius: 4px; border: 1px solid var(--line); background: var(--surface); }
    .stTextInput input:focus { border-color: var(--accent); box-shadow: 0 0 0 1px var(--accent); }
    [data-testid="stFileUploader"] section { border-radius: 4px; border: 1px dashed #B9C1C8; background: #FAFBFB; }

    /* tabs */
    .stTabs [data-baseweb="tab-list"] { gap: 1.75rem; border-bottom: 1px solid var(--line); }
    .stTabs [data-baseweb="tab"] { padding: 0.6rem 0; color: var(--muted); font-weight: 500; }
    .stTabs [aria-selected="true"] { color: var(--ink); }
    .stTabs [data-baseweb="tab-highlight"] { background: var(--accent); height: 2px; }
    .stTabs [data-baseweb="tab-border"] { display: none; }

    /* content panels */
    .panel { background: var(--surface); border: 1px solid var(--line); border-radius: 6px; padding: 1.5rem 1.75rem; line-height: 1.7; }
    .meta { color: var(--muted); font-size: 0.88rem; margin-bottom: 1.5rem; word-break: break-all; }
    .item { display: flex; gap: 0.9rem; padding: 0.85rem 0; border-bottom: 1px solid var(--line); line-height: 1.6; }
    .item:last-child { border-bottom: none; }
    .item .dot { flex: 0 0 6px; height: 6px; border-radius: 50%; background: var(--accent); margin-top: 0.65rem; }
    .empty { color: var(--muted); font-style: italic; }

    /* intro */
    .intro { margin: 3rem 0 1.75rem 0; }
    .intro p { color: var(--muted); line-height: 1.7; font-size: 1.02rem; max-width: 620px; }

    /* chat */
    [data-testid="stChatMessage"] { background: var(--surface); border: 1px solid var(--line); border-radius: 6px; padding: 1rem 1.25rem; }
    [data-testid="stChatInput"] textarea { border-radius: 4px; }

    :focus-visible { outline: 2px solid var(--accent); outline-offset: 2px; }

    /* ---- mobile ---- */
    .panel, .item, .meta { overflow-wrap: anywhere; }

    @media (max-width: 640px) {
        .block-container { padding: 1.25rem 1rem 3rem 1rem; }
        h1 { font-size: 1.6rem; }
        .intro { margin: 1.5rem 0 1rem 0; }
        .intro p { font-size: 0.95rem; }
        .panel { padding: 1rem 1.1rem; }
        .item { gap: 0.7rem; padding: 0.7rem 0; }
        .stTabs [data-baseweb="tab-list"] { gap: 1.1rem; overflow-x: auto; }
        .stTabs [data-baseweb="tab"] { white-space: nowrap; font-size: 0.92rem; }
        [data-testid="stChatMessage"] { padding: 0.75rem 0.85rem; }
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ---------------------------------------------------------------- helpers
STEPS = [
    "Preparing audio",
    "Transcribing speech",
    "Writing summary",
    "Extracting action items, decisions and questions",
    "Indexing transcript for chat",
]


def is_valid_url(text: str) -> bool:
    parsed = urlparse(text.strip())
    return parsed.scheme in ("http", "https") and bool(parsed.netloc)


def to_items(value) -> list:
    """Normalise LLM output (str or list) to a list of clean lines."""
    if not value:
        return []
    if isinstance(value, (list, tuple)):
        return [str(v).strip() for v in value if str(v).strip()]
    lines = []
    for line in str(value).splitlines():
        line = line.strip().lstrip("-•*").strip()
        if line[:2].rstrip(".)").isdigit():
            line = line.split(" ", 1)[-1].strip()
        if not line:
            continue
        # skip the model's intro sentence, e.g. "Here is a numbered list ...:"
        if line.lower().startswith(("here is", "here are")) or line.endswith(":"):
            continue
        lines.append(line)
    return lines


def render_items(value, empty_text: str):
    items = to_items(value)
    if not items:
        st.markdown(f'<div class="panel empty">{empty_text}</div>', unsafe_allow_html=True)
        return

    def fmt(text: str) -> str:
        safe = html.escape(text)
        safe = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", safe)
        return safe.replace("*", "")  # remove any leftover single asterisks

    rows = "".join(
        f'<div class="item"><span class="dot"></span><span>{fmt(i)}</span></div>'
        for i in items
    )
    st.markdown(f'<div class="panel">{rows}</div>', unsafe_allow_html=True)


def build_report(r: dict) -> str:
    def bullets(v):
        return "\n".join(f"- {i}" for i in to_items(v)) or "None identified."

    return (
        f"# {r['title']}\n\n"
        f"Generated {datetime.now():%d %b %Y, %H:%M}\n\n"
        f"## Summary\n\n{r['summary']}\n\n"
        f"## Action items\n\n{bullets(r['action_items'])}\n\n"
        f"## Key decisions\n\n{bullets(r['key_decisions'])}\n\n"
        f"## Open questions\n\n{bullets(r['open_questions'])}\n\n"
        f"## Transcript\n\n{r['transcript']}\n"
    )


def clean(text: str) -> str:
    """Strip Markdown bold/italic markers for plain-text output."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    return text.replace("*", "").strip()


def build_report_docx(r: dict) -> bytes:
    doc = Document()
    doc.add_heading(clean(r["title"]), level=0)
    doc.add_paragraph(f"Generated {datetime.now():%d %b %Y, %H:%M}")

    sections = [
        ("Summary", r["summary"]),
        ("Action items", r["action_items"]),
        ("Key decisions", r["key_decisions"]),
        ("Open questions", r["open_questions"]),
    ]
    for name, value in sections:
        doc.add_heading(name, level=1)
        items = to_items(value)
        if not items:
            doc.add_paragraph("None identified.")
        for item in items:
            doc.add_paragraph(clean(item), style="List Bullet")

    doc.add_heading("Transcript", level=1)
    doc.add_paragraph(r["transcript"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def run_pipeline_ui(source: str) -> dict:
    """Runs the same steps as the CLI, reporting progress in the UI."""
    with st.status("Analyzing your video", expanded=True) as status:
        st.write(STEPS[0])
        chunks = process_input(source)

        st.write(STEPS[1])
        try:
            transcript = transcribe_all(chunks)
        finally:
            cleanup(chunks)

        st.write(STEPS[2])
        title = generate_title(transcript)
        summary = summarize(transcript)

        st.write(STEPS[3])
        action_items = extract_action_items(transcript)
        decisions = extract_key_decisions(transcript)
        questions = extract_questions(transcript)

        st.write(STEPS[4])
        rag_chain = build_rag_chain(transcript)

        status.update(label="Analysis complete", state="complete", expanded=False)

    return {
        "title": title,
        "transcript": transcript,
        "summary": summary,
        "action_items": action_items,
        "key_decisions": decisions,
        "open_questions": questions,
        "rag_chain": rag_chain,
    }


# ---------------------------------------------------------------- state
st.session_state.setdefault("result", None)
st.session_state.setdefault("messages", [])
st.session_state.setdefault("source_label", "")

result = st.session_state.result

# ================================================================ INPUT VIEW
if not result:
    st.markdown('<h1 class="brand">ClipSum: Video Assistant</h1>', unsafe_allow_html=True)
    st.markdown(
        """
        <div class="intro">
            <p>Paste a video link or upload a file. You'll get a summary, action items,
            key decisions and open questions, and you can ask follow-up questions about
            anything that was said.</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    mode = st.radio(
        "Source",
        ["Upload a file", "Paste a link"],
        horizontal=True,
        label_visibility="collapsed",
    )

    url, upload = "", None
    if mode == "Paste a link":
        url = st.text_input(
            "Video link",
            placeholder="https://…",
            help="YouTube, Reddit, or any other page with a video that your downloader supports.",
        )
        if url.strip() and not is_valid_url(url):
            st.caption("Enter a full link starting with http:// or https://")
        st.caption("Some sites block downloads from cloud servers. If this fails, upload the file instead.")
    else:
        upload = st.file_uploader(
            "Audio or video file",
            type=["mp4", "mov", "mkv", "webm", "mp3", "wav", "m4a"],
        )

    ready = is_valid_url(url) if mode == "Paste a link" else upload is not None
    analyze = st.button("Analyze", type="primary", disabled=not ready)

    if analyze:
        tmp_path = None
        succeeded = False
        try:
            if mode == "Paste a link":
                source = url.strip()
                label = source
            else:
                suffix = os.path.splitext(upload.name)[1]
                with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                    tmp.write(upload.getbuffer())
                    tmp_path = tmp.name
                source = tmp_path
                label = upload.name

            st.session_state.result = run_pipeline_ui(source)
            st.session_state.source_label = label
            st.session_state.messages = []
            succeeded = True
        except Exception as e:
            st.error(f"The analysis could not be completed. {e}")
        finally:
            if tmp_path and os.path.exists(tmp_path):
                os.remove(tmp_path)

        if succeeded:
            st.rerun()

    st.stop()

# ================================================================ RESULTS VIEW
top_left, top_mid, top_right = st.columns([3, 1.4, 1.2])
with top_left:
    st.markdown('<h1 class="brand">ClipSum: Video Assistant</h1>', unsafe_allow_html=True)
with top_mid:
    st.download_button(
        "Download report",
        data=build_report_docx(result),
        file_name="video_report.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True,
    )
with top_right:
    if st.button("New video", use_container_width=True):
        st.session_state.result = None
        st.session_state.messages = []
        st.rerun()

st.title(result["title"])
st.markdown(
    f'<div class="meta">{html.escape(st.session_state.source_label)} &nbsp;|&nbsp; '
    f'{len(result["transcript"].split()):,} words transcribed</div>',
    unsafe_allow_html=True,
)

tab_sum, tab_act, tab_dec, tab_q, tab_tr, tab_chat = st.tabs(
    ["Summary", "Action items", "Decisions", "Open questions", "Transcript", "Ask a question"]
)

with tab_sum:
    with st.container(border=True):
        st.markdown(result["summary"])

with tab_act:
    render_items(result["action_items"], "No action items were identified.")

with tab_dec:
    render_items(result["key_decisions"], "No key decisions were identified.")

with tab_q:
    render_items(result["open_questions"], "No open questions were identified.")

with tab_tr:
    st.text_area(
        "Full transcript",
        value=result["transcript"],
        height=460,
        label_visibility="collapsed",
    )
    st.download_button("Download transcript", result["transcript"], file_name="transcript.txt")
    
with tab_chat:
    chat_box = st.container()          # messages go here, above the input
    question = st.chat_input("Ask about this video")

    with chat_box:
        if not st.session_state.messages and not question:
            st.markdown(
                '<p class="empty">Ask about anything discussed, for example “What was decided about the timeline?”</p>',
                unsafe_allow_html=True,
            )
        for m in st.session_state.messages:
            with st.chat_message(m["role"]):
                st.markdown(m["content"])

        if question:
            st.session_state.messages.append({"role": "user", "content": question})
            with st.chat_message("user"):
                st.markdown(question)
            with st.chat_message("assistant"):
                with st.spinner("Searching the transcript"):
                    try:
                        answer = ask_question(result["rag_chain"], question)
                    except Exception as e:
                        answer = f"I couldn't answer that. {e}"
                st.markdown(answer)
            st.session_state.messages.append({"role": "assistant", "content": answer})











# import os
# import html
# import re
# import tempfile
# from datetime import datetime
# from urllib.parse import urlparse

# import streamlit as st
# from dotenv import load_dotenv

# from utils.audio_processor import process_input, cleanup
# from core.transcriber import transcribe_all
# from core.summarizer import summarize, generate_title
# from core.extractor import extract_action_items, extract_key_decisions, extract_questions
# from core.rag_engine import build_rag_chain, ask_question

# load_dotenv()

# st.set_page_config(
#     page_title="Video Assistant",
#     page_icon="◼",
#     layout="centered",
#     initial_sidebar_state="collapsed",
# )

# # ---------------------------------------------------------------- styling
# st.markdown(
#     """
#     <style>
#     @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600&family=Source+Serif+4:wght@500;600&display=swap');

#     :root {
#         --ink: #1F2933;
#         --muted: #667380;
#         --line: #DADFE3;
#         --surface: #FFFFFF;
#         --accent: #3B4A5A;
#     }

#     html, body, [class*="css"], .stMarkdown, .stTextInput, .stButton, .stTabs, .stRadio {
#         font-family: 'Inter', -apple-system, 'Segoe UI', sans-serif;
#         color: var(--ink);
#     }

#     /* this app has no sidebar */
#     section[data-testid="stSidebar"],
#     [data-testid="stSidebarCollapsedControl"],
#     [data-testid="collapsedControl"],
#     [data-testid="stExpandSidebarButton"] { display: none !important; }

#     #MainMenu, footer, header[data-testid="stHeader"] { visibility: hidden; height: 0; }
#     .block-container { padding-top: 2.5rem; padding-bottom: 4rem; max-width: 900px; }

#     h1, h2, h3 { font-family: 'Source Serif 4', Georgia, serif; font-weight: 600; letter-spacing: -0.01em; }
#     h1 { font-size: 2.1rem; line-height: 1.25; }

#     .brand { font-family: 'Source Serif 4', Georgia, serif; font-size: 1.3rem; font-weight: 600; margin: 0; }
#     .brand-sub { color: var(--muted); font-size: 0.9rem; margin: 0.15rem 0 0 0; }

#     /* buttons */
#     .stButton > button, .stDownloadButton > button {
#         border-radius: 4px; border: 1px solid var(--line); background: var(--surface);
#         color: var(--ink); font-weight: 500; padding: 0.55rem 1rem; transition: border-color .15s, background .15s;
#     }
#     .stButton > button:hover, .stDownloadButton > button:hover { border-color: var(--accent); color: var(--accent); background: var(--surface); }
#     .stButton > button[kind="primary"] { background: var(--accent); border-color: var(--accent); color: #fff; }
#     .stButton > button[kind="primary"]:hover { background: #2C3947; border-color: #2C3947; color: #fff; }
#     .stButton > button:disabled { opacity: 0.5; }

#     /* inputs */
#     .stTextInput input { border-radius: 4px; border: 1px solid var(--line); background: var(--surface); }
#     .stTextInput input:focus { border-color: var(--accent); box-shadow: 0 0 0 1px var(--accent); }
#     [data-testid="stFileUploader"] section { border-radius: 4px; border: 1px dashed #B9C1C8; background: #FAFBFB; }

#     /* tabs */
#     .stTabs [data-baseweb="tab-list"] { gap: 1.75rem; border-bottom: 1px solid var(--line); }
#     .stTabs [data-baseweb="tab"] { padding: 0.6rem 0; color: var(--muted); font-weight: 500; }
#     .stTabs [aria-selected="true"] { color: var(--ink); }
#     .stTabs [data-baseweb="tab-highlight"] { background: var(--accent); height: 2px; }
#     .stTabs [data-baseweb="tab-border"] { display: none; }

#     /* content panels */
#     .panel { background: var(--surface); border: 1px solid var(--line); border-radius: 6px; padding: 1.5rem 1.75rem; line-height: 1.7; }
#     .meta { color: var(--muted); font-size: 0.88rem; margin-bottom: 1.5rem; word-break: break-all; }
#     .item { display: flex; gap: 0.9rem; padding: 0.85rem 0; border-bottom: 1px solid var(--line); line-height: 1.6; }
#     .item:last-child { border-bottom: none; }
#     .item .dot { flex: 0 0 6px; height: 6px; border-radius: 50%; background: var(--accent); margin-top: 0.65rem; }
#     .empty { color: var(--muted); font-style: italic; }

#     /* intro */
#     .intro { margin: 3rem 0 1.75rem 0; }
#     .intro p { color: var(--muted); line-height: 1.7; font-size: 1.02rem; max-width: 620px; }

#     /* chat */
#     [data-testid="stChatMessage"] { background: var(--surface); border: 1px solid var(--line); border-radius: 6px; padding: 1rem 1.25rem; }
#     [data-testid="stChatInput"] textarea { border-radius: 4px; }

#     :focus-visible { outline: 2px solid var(--accent); outline-offset: 2px; }
#     </style>
#     """,
#     unsafe_allow_html=True,
# )

# # ---------------------------------------------------------------- helpers
# STEPS = [
#     "Preparing audio",
#     "Transcribing speech",
#     "Writing summary",
#     "Extracting action items, decisions and questions",
#     "Indexing transcript for chat",
# ]


# def is_valid_url(text: str) -> bool:
#     parsed = urlparse(text.strip())
#     return parsed.scheme in ("http", "https") and bool(parsed.netloc)


# def to_items(value) -> list:
#     """Normalise LLM output (str or list) to a list of clean lines."""
#     if not value:
#         return []
#     if isinstance(value, (list, tuple)):
#         return [str(v).strip() for v in value if str(v).strip()]
#     lines = []
#     for line in str(value).splitlines():
#         line = line.strip().lstrip("-•*").strip()
#         if line[:2].rstrip(".)").isdigit():
#             line = line.split(" ", 1)[-1].strip()
#         if not line:
#             continue
#         # skip the model's intro sentence, e.g. "Here is a numbered list ...:"
#         if line.lower().startswith(("here is", "here are")) or line.endswith(":"):
#             continue
#         lines.append(line)
#     return lines


# def render_items(value, empty_text: str):
#     items = to_items(value)
#     if not items:
#         st.markdown(f'<div class="panel empty">{empty_text}</div>', unsafe_allow_html=True)
#         return

#     def fmt(text: str) -> str:
#         safe = html.escape(text)
#         safe = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", safe)
#         return safe.replace("*", "")  # remove any leftover single asterisks

#     rows = "".join(
#         f'<div class="item"><span class="dot"></span><span>{fmt(i)}</span></div>'
#         for i in items
#     )
#     st.markdown(f'<div class="panel">{rows}</div>', unsafe_allow_html=True)


# def build_report(r: dict) -> str:
#     def bullets(v):
#         return "\n".join(f"- {i}" for i in to_items(v)) or "None identified."

#     return (
#         f"# {r['title']}\n\n"
#         f"Generated {datetime.now():%d %b %Y, %H:%M}\n\n"
#         f"## Summary\n\n{r['summary']}\n\n"
#         f"## Action items\n\n{bullets(r['action_items'])}\n\n"
#         f"## Key decisions\n\n{bullets(r['key_decisions'])}\n\n"
#         f"## Open questions\n\n{bullets(r['open_questions'])}\n\n"
#         f"## Transcript\n\n{r['transcript']}\n"
#     )


# def run_pipeline_ui(source: str) -> dict:
#     """Runs the same steps as the CLI, reporting progress in the UI."""
#     with st.status("Analyzing your video", expanded=True) as status:
#         st.write(STEPS[0])
#         chunks = process_input(source)

#         st.write(STEPS[1])
#         try:
#             transcript = transcribe_all(chunks)
#         finally:
#             cleanup(chunks)

#         st.write(STEPS[2])
#         title = generate_title(transcript)
#         summary = summarize(transcript)

#         st.write(STEPS[3])
#         action_items = extract_action_items(transcript)
#         decisions = extract_key_decisions(transcript)
#         questions = extract_questions(transcript)

#         st.write(STEPS[4])
#         rag_chain = build_rag_chain(transcript)

#         status.update(label="Analysis complete", state="complete", expanded=False)

#     return {
#         "title": title,
#         "transcript": transcript,
#         "summary": summary,
#         "action_items": action_items,
#         "key_decisions": decisions,
#         "open_questions": questions,
#         "rag_chain": rag_chain,
#     }


# # ---------------------------------------------------------------- state
# st.session_state.setdefault("result", None)
# st.session_state.setdefault("messages", [])
# st.session_state.setdefault("source_label", "")

# result = st.session_state.result

# # ================================================================ INPUT VIEW
# if not result:
#     st.markdown('<h1 class="brand">Video Assistant</h1>', unsafe_allow_html=True)
#     st.markdown(
#         """
#         <div class="intro">
#             <p>Paste a video link or upload a file. You'll get a summary, action items,
#             key decisions and open questions, and you can ask follow-up questions about
#             anything that was said.</p>
#         </div>
#         """,
#         unsafe_allow_html=True,
#     )

#     mode = st.radio(
#         "Source",
#         ["Upload a file", "Paste a link"],
#         horizontal=True,
#         label_visibility="collapsed",
#     )

#     url, upload = "", None
#     if mode == "Paste a link":
#         url = st.text_input(
#             "Video link",
#             placeholder="https://…",
#             help="YouTube, Reddit, or any other page with a video that your downloader supports.",
#         )
#         if url.strip() and not is_valid_url(url):
#             st.caption("Enter a full link starting with http:// or https://")
#         st.caption("Some sites block downloads from cloud servers. If this fails, upload the file instead.")
#     else:
#         upload = st.file_uploader(
#             "Audio or video file",
#             type=["mp4", "mov", "mkv", "webm", "mp3", "wav", "m4a"],
#         )

#     ready = is_valid_url(url) if mode == "Paste a link" else upload is not None
#     analyze = st.button("Analyze", type="primary", disabled=not ready)

#     if analyze:
#         tmp_path = None
#         succeeded = False
#         try:
#             if mode == "Paste a link":
#                 source = url.strip()
#                 label = source
#             else:
#                 suffix = os.path.splitext(upload.name)[1]
#                 with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
#                     tmp.write(upload.getbuffer())
#                     tmp_path = tmp.name
#                 source = tmp_path
#                 label = upload.name

#             st.session_state.result = run_pipeline_ui(source)
#             st.session_state.source_label = label
#             st.session_state.messages = []
#             succeeded = True
#         except Exception as e:
#             st.error(f"The analysis could not be completed. {e}")
#         finally:
#             if tmp_path and os.path.exists(tmp_path):
#                 os.remove(tmp_path)

#         if succeeded:
#             st.rerun()

#     st.stop()

# # ================================================================ RESULTS VIEW
# top_left, top_mid, top_right = st.columns([3, 1.4, 1.2])
# with top_left:
#     st.markdown('<h1 class="brand">Video Assistant</h1>', unsafe_allow_html=True)
# with top_mid:
#     st.download_button(
#         "Download report",
#         data=build_report(result),
#         file_name="video_report.md",
#         mime="text/markdown",
#         use_container_width=True,
#     )
# with top_right:
#     if st.button("New video", use_container_width=True):
#         st.session_state.result = None
#         st.session_state.messages = []
#         st.rerun()

# st.title(result["title"])
# st.markdown(
#     f'<div class="meta">{html.escape(st.session_state.source_label)} &nbsp;|&nbsp; '
#     f'{len(result["transcript"].split()):,} words transcribed</div>',
#     unsafe_allow_html=True,
# )

# tab_sum, tab_act, tab_dec, tab_q, tab_tr, tab_chat = st.tabs(
#     ["Summary", "Action items", "Decisions", "Open questions", "Transcript", "Ask a question"]
# )

# with tab_sum:
#     with st.container(border=True):
#         st.markdown(result["summary"])

# with tab_act:
#     render_items(result["action_items"], "No action items were identified.")

# with tab_dec:
#     render_items(result["key_decisions"], "No key decisions were identified.")

# with tab_q:
#     render_items(result["open_questions"], "No open questions were identified.")

# with tab_tr:
#     st.text_area(
#         "Full transcript",
#         value=result["transcript"],
#         height=460,
#         label_visibility="collapsed",
#     )
#     st.download_button("Download transcript", result["transcript"], file_name="transcript.txt")

# with tab_chat:
#     if not st.session_state.messages:
#         st.markdown(
#             '<p class="empty">Ask about anything discussed, for example “What was decided about the timeline?”</p>',
#             unsafe_allow_html=True,
#         )
#     for m in st.session_state.messages:
#         with st.chat_message(m["role"]):
#             st.markdown(m["content"])

#     if question := st.chat_input("Ask about this video"):
#         st.session_state.messages.append({"role": "user", "content": question})
#         with st.chat_message("user"):
#             st.markdown(question)
#         with st.chat_message("assistant"):
#             with st.spinner("Searching the transcript"):
#                 try:
#                     answer = ask_question(result["rag_chain"], question)
#                 except Exception as e:
#                     answer = f"I couldn't answer that. {e}"
#             st.markdown(answer)
#         st.session_state.messages.append({"role": "assistant", "content": answer})


# import os
# import html
# import tempfile
# from datetime import datetime
# from urllib.parse import urlparse

# import streamlit as st
# from dotenv import load_dotenv

# from utils.audio_processor import process_input
# from core.transcriber import transcribe_all
# from core.summarizer import summarize, generate_title
# from core.extractor import extract_action_items, extract_key_decisions, extract_questions
# from core.rag_engine import build_rag_chain, ask_question

# load_dotenv()

# st.set_page_config(
#     page_title="Video Assistant",
#     page_icon="◼",
#     layout="centered",
#     initial_sidebar_state="collapsed",
# )

# # ---------------------------------------------------------------- styling
# st.markdown(
#     """
#     <style>
#     @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600&family=Source+Serif+4:wght@500;600&display=swap');

#     :root {
#         --ink: #1F2933;
#         --muted: #667380;
#         --line: #DADFE3;
#         --surface: #FFFFFF;
#         --accent: #3B4A5A;
#     }

#     html, body, [class*="css"], .stMarkdown, .stTextInput, .stButton, .stTabs, .stRadio {
#         font-family: 'Inter', -apple-system, 'Segoe UI', sans-serif;
#         color: var(--ink);
#     }

#     /* this app has no sidebar */
#     section[data-testid="stSidebar"],
#     [data-testid="stSidebarCollapsedControl"],
#     [data-testid="collapsedControl"],
#     [data-testid="stExpandSidebarButton"] { display: none !important; }

#     #MainMenu, footer, header[data-testid="stHeader"] { visibility: hidden; height: 0; }
#     .block-container { padding-top: 2.5rem; padding-bottom: 4rem; max-width: 900px; }

#     h1, h2, h3 { font-family: 'Source Serif 4', Georgia, serif; font-weight: 600; letter-spacing: -0.01em; }
#     h1 { font-size: 2.1rem; line-height: 1.25; }

#     .brand { font-family: 'Source Serif 4', Georgia, serif; font-size: 1.3rem; font-weight: 600; margin: 0; }
#     .brand-sub { color: var(--muted); font-size: 0.9rem; margin: 0.15rem 0 0 0; }

#     /* buttons */
#     .stButton > button, .stDownloadButton > button {
#         border-radius: 4px; border: 1px solid var(--line); background: var(--surface);
#         color: var(--ink); font-weight: 500; padding: 0.55rem 1rem; transition: border-color .15s, background .15s;
#     }
#     .stButton > button:hover, .stDownloadButton > button:hover { border-color: var(--accent); color: var(--accent); background: var(--surface); }
#     .stButton > button[kind="primary"] { background: var(--accent); border-color: var(--accent); color: #fff; }
#     .stButton > button[kind="primary"]:hover { background: #2C3947; border-color: #2C3947; color: #fff; }
#     .stButton > button:disabled { opacity: 0.5; }

#     /* inputs */
#     .stTextInput input { border-radius: 4px; border: 1px solid var(--line); background: var(--surface); }
#     .stTextInput input:focus { border-color: var(--accent); box-shadow: 0 0 0 1px var(--accent); }
#     [data-testid="stFileUploader"] section { border-radius: 4px; border: 1px dashed #B9C1C8; background: #FAFBFB; }

#     /* tabs */
#     .stTabs [data-baseweb="tab-list"] { gap: 1.75rem; border-bottom: 1px solid var(--line); }
#     .stTabs [data-baseweb="tab"] { padding: 0.6rem 0; color: var(--muted); font-weight: 500; }
#     .stTabs [aria-selected="true"] { color: var(--ink); }
#     .stTabs [data-baseweb="tab-highlight"] { background: var(--accent); height: 2px; }
#     .stTabs [data-baseweb="tab-border"] { display: none; }

#     /* content panels */
#     .panel { background: var(--surface); border: 1px solid var(--line); border-radius: 6px; padding: 1.5rem 1.75rem; line-height: 1.7; }
#     .meta { color: var(--muted); font-size: 0.88rem; margin-bottom: 1.5rem; word-break: break-all; }
#     .item { display: flex; gap: 0.9rem; padding: 0.85rem 0; border-bottom: 1px solid var(--line); line-height: 1.6; }
#     .item:last-child { border-bottom: none; }
#     .item .dot { flex: 0 0 6px; height: 6px; border-radius: 50%; background: var(--accent); margin-top: 0.65rem; }
#     .empty { color: var(--muted); font-style: italic; }

#     /* intro */
#     .intro { margin: 3rem 0 1.75rem 0; }
#     .intro p { color: var(--muted); line-height: 1.7; font-size: 1.02rem; max-width: 620px; }

#     /* chat */
#     [data-testid="stChatMessage"] { background: var(--surface); border: 1px solid var(--line); border-radius: 6px; padding: 1rem 1.25rem; }
#     [data-testid="stChatInput"] textarea { border-radius: 4px; }

#     :focus-visible { outline: 2px solid var(--accent); outline-offset: 2px; }
#     </style>
#     """,
#     unsafe_allow_html=True,
# )

# # ---------------------------------------------------------------- helpers
# STEPS = [
#     "Preparing audio",
#     "Transcribing speech",
#     "Writing summary",
#     "Extracting action items, decisions and questions",
#     "Indexing transcript for chat",
# ]


# def is_valid_url(text: str) -> bool:
#     parsed = urlparse(text.strip())
#     return parsed.scheme in ("http", "https") and bool(parsed.netloc)


# def to_items(value) -> list:
#     """Normalise LLM output (str or list) to a list of clean lines."""
#     if not value:
#         return []
#     if isinstance(value, (list, tuple)):
#         return [str(v).strip() for v in value if str(v).strip()]
#     lines = []
#     for line in str(value).splitlines():
#         line = line.strip().lstrip("-•*").strip()
#         if line[:2].rstrip(".)").isdigit():
#             line = line.split(" ", 1)[-1].strip()
#         if line:
#             lines.append(line)
#     return lines


# def render_items(value, empty_text: str):
#     items = to_items(value)
#     if not items:
#         st.markdown(f'<div class="panel empty">{empty_text}</div>', unsafe_allow_html=True)
#         return
#     rows = "".join(f'<div class="item"><span class="dot"></span><span>{i}</span></div>' for i in items)
#     st.markdown(f'<div class="panel">{rows}</div>', unsafe_allow_html=True)


# def build_report(r: dict) -> str:
#     def bullets(v):
#         return "\n".join(f"- {i}" for i in to_items(v)) or "None identified."

#     return (
#         f"# {r['title']}\n\n"
#         f"Generated {datetime.now():%d %b %Y, %H:%M}\n\n"
#         f"## Summary\n\n{r['summary']}\n\n"
#         f"## Action items\n\n{bullets(r['action_items'])}\n\n"
#         f"## Key decisions\n\n{bullets(r['key_decisions'])}\n\n"
#         f"## Open questions\n\n{bullets(r['open_questions'])}\n\n"
#         f"## Transcript\n\n{r['transcript']}\n"
#     )


# def run_pipeline_ui(source: str) -> dict:
#     """Runs the same steps as the CLI, reporting progress in the UI."""
#     with st.status("Analyzing your video", expanded=True) as status:
#         st.write(STEPS[0])
#         chunks = process_input(source)

#         st.write(STEPS[1])
#         transcript = transcribe_all(chunks)

#         st.write(STEPS[2])
#         title = generate_title(transcript)
#         summary = summarize(transcript)

#         st.write(STEPS[3])
#         action_items = extract_action_items(transcript)
#         decisions = extract_key_decisions(transcript)
#         questions = extract_questions(transcript)

#         st.write(STEPS[4])
#         rag_chain = build_rag_chain(transcript)

#         status.update(label="Analysis complete", state="complete", expanded=False)

#     return {
#         "title": title,
#         "transcript": transcript,
#         "summary": summary,
#         "action_items": action_items,
#         "key_decisions": decisions,
#         "open_questions": questions,
#         "rag_chain": rag_chain,
#     }


# # ---------------------------------------------------------------- state
# st.session_state.setdefault("result", None)
# st.session_state.setdefault("messages", [])
# st.session_state.setdefault("source_label", "")

# result = st.session_state.result

# # ================================================================ INPUT VIEW
# if not result:
#     st.markdown('<p class="brand">Video Assistant</p>', unsafe_allow_html=True)
#     st.markdown(
#         """
#         <div class="intro">
#             <h1>Turn a recording into something you can act on.</h1>
#             <p>Paste a video link or upload a file. You'll get a summary, action items,
#             key decisions and open questions, and you can ask follow-up questions about
#             anything that was said.</p>
#         </div>
#         """,
#         unsafe_allow_html=True,
#     )

#     mode = st.radio(
#         "Source",
#         ["Paste a link", "Upload a file"],
#         horizontal=True,
#         label_visibility="collapsed",
#     )

#     url, upload = "", None
#     if mode == "Paste a link":
#         url = st.text_input(
#             "Video link",
#             placeholder="https://…",
#             help="YouTube, Reddit, or any other page with a video that your downloader supports.",
#         )
#         if url.strip() and not is_valid_url(url):
#             st.caption("Enter a full link starting with http:// or https://")
#     else:
#         upload = st.file_uploader(
#             "Audio or video file",
#             type=["mp4", "mov", "mkv", "webm", "mp3", "wav", "m4a"],
#         )

#     ready = is_valid_url(url) if mode == "Paste a link" else upload is not None
#     analyze = st.button("Analyze", type="primary", disabled=not ready)

#     if analyze:
#         tmp_path = None
#         succeeded = False
#         try:
#             if mode == "Paste a link":
#                 source = url.strip()
#                 label = source
#             else:
#                 suffix = os.path.splitext(upload.name)[1]
#                 with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
#                     tmp.write(upload.getbuffer())
#                     tmp_path = tmp.name
#                 source = tmp_path
#                 label = upload.name

#             st.session_state.result = run_pipeline_ui(source)
#             st.session_state.source_label = label
#             st.session_state.messages = []
#             succeeded = True
#         except Exception as e:
#             st.error(f"The analysis could not be completed. {e}")
#         finally:
#             if tmp_path and os.path.exists(tmp_path):
#                 os.remove(tmp_path)

#         if succeeded:
#             st.rerun()

#     st.stop()

# # ================================================================ RESULTS VIEW
# top_left, top_mid, top_right = st.columns([3, 1.4, 1.2])
# with top_left:
#     st.markdown('<p class="brand">Video Assistant</p>', unsafe_allow_html=True)
# with top_mid:
#     st.download_button(
#         "Download report",
#         data=build_report(result),
#         file_name="video_report.md",
#         mime="text/markdown",
#         use_container_width=True,
#     )
# with top_right:
#     if st.button("New video", use_container_width=True):
#         st.session_state.result = None
#         st.session_state.messages = []
#         st.rerun()

# st.title(result["title"])
# st.markdown(
#     f'<div class="meta">{st.session_state.source_label} &nbsp;|&nbsp; '
#     f'{len(result["transcript"].split()):,} words transcribed</div>',
#     unsafe_allow_html=True,
# )

# tab_sum, tab_act, tab_dec, tab_q, tab_tr, tab_chat = st.tabs(
#     ["Summary", "Action items", "Decisions", "Open questions", "Transcript", "Ask a question"]
# )

# with tab_sum:
#     st.markdown(f'<div class="panel">{result["summary"]}</div>', unsafe_allow_html=True)

# with tab_act:
#     render_items(result["action_items"], "No action items were identified.")

# with tab_dec:
#     render_items(result["key_decisions"], "No key decisions were identified.")

# with tab_q:
#     render_items(result["open_questions"], "No open questions were identified.")

# with tab_tr:
#     st.text_area(
#         "Full transcript",
#         value=result["transcript"],
#         height=460,
#         label_visibility="collapsed",
#     )
#     st.download_button("Download transcript", result["transcript"], file_name="transcript.txt")

# with tab_chat:
#     if not st.session_state.messages:
#         st.markdown(
#             '<p class="empty">Ask about anything discussed, for example “What was decided about the timeline?”</p>',
#             unsafe_allow_html=True,
#         )
#     for m in st.session_state.messages:
#         with st.chat_message(m["role"]):
#             st.markdown(m["content"])

#     if question := st.chat_input("Ask about this video"):
#         st.session_state.messages.append({"role": "user", "content": question})
#         with st.chat_message("user"):
#             st.markdown(question)
#         with st.chat_message("assistant"):
#             with st.spinner("Searching the transcript"):
#                 try:
#                     answer = ask_question(result["rag_chain"], question)
#                 except Exception as e:
#                     answer = f"I couldn't answer that. {e}"
#             st.markdown(answer)
#         st.session_state.messages.append({"role": "assistant", "content": answer})